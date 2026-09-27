"""Reading text out of a PDF or DOCX.

The document is the most sensitive artefact this product handles — it states
everything feature 001 spends effort protecting, and often an address and a
previous salary besides. FR-004 says it is not retained.

That promise is kept by never writing it. Bytes arrive, text comes out, the
bytes are released when this function returns. There is no temporary file, no
object store, and therefore no cleanup step that can be skipped, no failure path
that leaves a resume at rest, and nothing to audit. A design that stored the
document and deleted it afterwards would be a promise; this is a property.

Extraction is deterministic library work, not a model task. Using a model here
would add a fabrication surface where none is needed, and would forfeit the
exact source text the evidence verifier depends on (research.md R-005).
"""

import enum
import io
import logging

logger = logging.getLogger(__name__)

# Below this, a document has no usable text layer. A scanned page yields a
# handful of stray characters rather than nothing at all, so a threshold
# distinguishes it from a genuinely short document more reliably than a zero
# check would.
MINIMUM_TEXT_LENGTH = 50

# An ordinary resume is a few pages. Beyond this a document is either not a
# resume or needs handling this feature does not offer, and the user is told so
# rather than waiting on something that will not finish well.
MAXIMUM_BYTES = 10 * 1024 * 1024


class DocumentProblem(enum.StrEnum):
    unreadable = "unreadable"
    unsupported = "unsupported"
    too_large = "too_large"


class DocumentError(Exception):
    """A document that cannot be turned into text, and why."""

    def __init__(self, problem: DocumentProblem, message: str) -> None:
        super().__init__(message)
        self.problem = problem
        self.message = message


PDF_MAGIC = b"%PDF-"
DOCX_MAGIC = b"PK\x03\x04"

CONTENT_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
}


def detect_kind(data: bytes, filename: str | None, content_type: str | None) -> str:
    """Decide what a document is from its contents, not its name.

    A filename can say anything. The leading bytes cannot, so they decide, and
    the declared type and extension are only consulted when the magic number is
    inconclusive.
    """
    if data.startswith(PDF_MAGIC):
        return "pdf"
    if data.startswith(DOCX_MAGIC):
        # A DOCX is a zip. So are several things that are not documents, which
        # python-docx will reject when it tries to open them.
        return "docx"

    if content_type in CONTENT_TYPES:
        return CONTENT_TYPES[content_type]

    suffix = (filename or "").lower().rsplit(".", 1)
    if len(suffix) == 2 and suffix[1] in ("pdf", "docx"):
        return suffix[1]

    raise DocumentError(
        DocumentProblem.unsupported,
        "That is not a PDF or a Word document. Those are the formats we can read.",
    )


def extract_text(data: bytes, filename: str | None = None, content_type: str | None = None) -> str:
    """Return the text of a document.

    The bytes are not written anywhere and are not retained beyond this call.
    """
    if not data:
        raise DocumentError(DocumentProblem.unreadable, "The file is empty.")

    if len(data) > MAXIMUM_BYTES:
        raise DocumentError(
            DocumentProblem.too_large,
            "That document is larger than we can read. Resumes are usually a few pages.",
        )

    kind = detect_kind(data, filename, content_type)
    text = _read_pdf(data) if kind == "pdf" else _read_docx(data)

    if len(text.strip()) < MINIMUM_TEXT_LENGTH:
        # Almost always a scan or a photograph. Said plainly, because no amount
        # of retrying will help and the user needs to know that.
        raise DocumentError(
            DocumentProblem.unreadable,
            "We could not find any text in that document. It may be a scan or a "
            "photograph rather than a document with text in it.",
        )

    logger.info("document read", extra={"entry_type": kind, "count": len(text)})
    return text


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentError(
                DocumentProblem.unreadable,
                "That document is password protected, so we cannot read it.",
            )
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except DocumentError:
        raise
    except (PdfReadError, ValueError, OSError) as error:
        raise DocumentError(
            DocumentProblem.unreadable, "That PDF could not be opened. It may be damaged."
        ) from error


def _read_docx(data: bytes) -> str:
    import zipfile

    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (
        PackageNotFoundError,
        # A DOCX is a zip, and a damaged one raises from zipfile rather than
        # from python-docx. BadZipFile is not an OSError, so it needs naming.
        zipfile.BadZipFile,
        KeyError,
        ValueError,
        OSError,
    ) as error:
        raise DocumentError(
            DocumentProblem.unreadable, "That Word document could not be opened."
        ) from error

    parts = [paragraph.text for paragraph in document.paragraphs]

    # Japanese resumes routinely put the substance in tables, so a paragraph-only
    # read would silently lose most of a 職務経歴書.
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)

    return "\n".join(part for part in parts if part.strip())
